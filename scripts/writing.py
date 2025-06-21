import openpyxl
import os
from scripts.scraping import get_combined_game_info  

def write_game_info_to_excel():
    """
    Reads from a text file and adds new game entries to the worksheet if they don't already exist.
    Calls scraping functions to retrieve game information and appends it to the worksheet.

    """
    # Define the platform for the games
    platform = "N64"

    # Read game titles from test.txt
    with open("test.txt", "r", encoding = "utf-8") as file:
        game_titles = [line.strip() for line in file.readlines() if line.strip()]

    # Check if the Excel file already exists
    file_exists = os.path.isfile("games.xlsx")

    # Open the workbook if it exists; otherwise, create a new one
    if file_exists:
        wb = openpyxl.load_workbook("games.xlsx")
        ws = wb.active
        # Collect existing titles to avoid duplicates
        existing_titles = {ws.cell(row=row, column=2).value for row in range(2, ws.max_row + 1)}
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Game Information"
        # Write headers if creating a new workbook
        headers = ["Platform", "Title", "Alt-Titles", "Publisher", "Developer", "Director", "Artist", "Writer", "Composer", "Genre", "Aggregate Score", "Theme", "Franchise"]
        ws.append(headers)
        existing_titles = set()  # No existing titles in a new file

    # Loop through each game, retrieve its information, and append it if it's new
    for title in game_titles:
        if title.split("(", 1)[0].strip() in existing_titles:
            # Skip if the title is already in the file
            print(f"Skipping {title}: already exists in the file.")
            continue

        print(f"\nStarting scraping for: {title}")
        game_info = get_combined_game_info(title)  # Get combined game info from both Wikipedia and Giant Bomb
        # Prepare the row data
        row = [
            platform,
            game_info.get("Title", ""),
            game_info.get("Alt-Titles", ""),
            game_info.get("Publisher", ""),
            game_info.get("Developer", ""),
            game_info.get("Director", ""),
            game_info.get("Artist", ""),
            game_info.get("Writer", ""),
            game_info.get("Composer", ""),
            game_info.get("Genre", ""),
            game_info.get("Aggregate Score", ""),
            game_info.get("Theme", ""),
            game_info.get("Franchise", "")
        ]
        
        # Append the row to the sheet
        ws.append(row)
        print(f"Added {title} to the game database")
        existing_titles.add(title)  # Add title to the set to track newly added games
        # Save the workbook
        wb.save("games.xlsx")

    print("Finished writing to the worksheet")